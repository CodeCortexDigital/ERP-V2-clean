"""Spreadsheet import (P10): templates, a preview that checks every row, the import itself, and a history.
/api/v1/auth/imports/"""
from __future__ import annotations

import csv
import io

from django.db import transaction
from django.http import HttpResponse
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from services.core.accounts.permissions import IsSchoolAdmin

from .models import ImportRun
from .specs import KINDS, norm, text

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024


def _kind(kind):
    spec = KINDS.get(kind)
    if spec is None:
        return None, Response({'error': 'Unknown import.'}, status=status.HTTP_404_NOT_FOUND)
    return spec, None


def _describe(kind, spec):
    return {'kind': kind, 'label': spec['label'], 'help': spec['help'], 'order': spec['order'],
            'columns': [{'key': k, 'label': label, 'required': req} for k, label, req, _ in spec['columns']]}


@api_view(['GET'])
@permission_classes([IsSchoolAdmin])
def kinds(request):
    runs = {r.kind: r for r in reversed(list(ImportRun.objects.all()[:200]))}
    out = []
    for kind, spec in sorted(KINDS.items(), key=lambda kv: kv[1]['order']):
        d = _describe(kind, spec)
        last = runs.get(kind)
        d['last_import'] = {'when': last.created_at, 'added': last.added} if last else None
        out.append(d)
    return Response({'kinds': out, 'max_rows': MAX_ROWS})


@api_view(['GET'])
@permission_classes([IsSchoolAdmin])
def template(request, kind):
    spec, error = _kind(kind)
    if error:
        return error
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="import-{kind}.csv"'
    response.write('﻿')  # Excel opens UTF-8 CSVs correctly with a BOM
    writer = csv.writer(response)
    writer.writerow([label + (' *' if req else '') for _, label, req, _ in spec['columns']])
    for row in spec['example']:
        writer.writerow(row)
    return response


def _read_rows(upload):
    """Rows as lists of cell values, from a CSV or an Excel file."""
    name = (upload.name or '').lower()
    if upload.size > MAX_BYTES:
        raise ValueError('The file is larger than 5 MB.')
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook

        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith('.xls'):
        raise ValueError('Old .xls files are not supported. In Excel choose File → Save As → .xlsx or CSV.')
    raw = upload.read()
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            content = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = content[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(content), dialect)]


def _map_headers(header, spec):
    """Which column of the file feeds which field. Returns (mapping, unknown_headers, missing_required)."""
    lookup = {}
    for key, label, _req, aliases in spec['columns']:
        for name in [key, label, *aliases]:
            lookup.setdefault(norm(name.replace('*', '')), key)
    mapping, unknown = {}, []
    for i, h in enumerate(header):
        k = lookup.get(norm(text(h).replace('*', '')))
        if k and k not in mapping.values():
            mapping[i] = k
        elif text(h):
            unknown.append(text(h))
    missing = [label for key, label, req, _ in spec['columns'] if req and key not in mapping.values()]
    return mapping, unknown, missing


def _check(request, kind):
    spec, error = _kind(kind)
    if error:
        return None, error
    upload = request.FILES.get('file')
    if upload is None:
        return None, Response({'error': 'Choose a CSV or Excel file.'}, status=status.HTTP_400_BAD_REQUEST)
    try:
        rows = _read_rows(upload)
    except ValueError as exc:
        return None, Response({'error': str(exc)}, status=status.HTTP_400_BAD_REQUEST)
    except Exception:
        return None, Response({'error': 'The file could not be read. Save it as CSV or .xlsx and try again.'}, status=status.HTTP_400_BAD_REQUEST)
    rows = [r for r in rows if any(text(c) for c in r)]
    if not rows:
        return None, Response({'error': 'The file is empty.'}, status=status.HTTP_400_BAD_REQUEST)
    mapping, unknown, missing = _map_headers(rows[0], spec)
    if missing:
        return None, Response({'error': f'Missing column{"s" if len(missing) > 1 else ""}: {", ".join(missing)}. Download the template to see the headings.',
                               'unknown_columns': unknown}, status=status.HTTP_400_BAD_REQUEST)
    body = rows[1:]
    if len(body) > MAX_ROWS:
        return None, Response({'error': f'Up to {MAX_ROWS} rows at a time; this file has {len(body)}. Split it into smaller files.'},
                              status=status.HTTP_400_BAD_REQUEST)
    ctx = {'school': request.tenant}
    checked = []
    for n, raw in enumerate(body, start=2):
        row = {key: (raw[i] if i < len(raw) else None) for i, key in mapping.items()}
        clean, errors, duplicate, note = spec['check'](row, ctx)
        state = 'error' if errors else 'duplicate' if duplicate else 'ready'
        checked.append({'row': n, 'state': state, 'messages': errors or ([duplicate] if duplicate else []), 'note': note,
                        'values': {k: text(v) for k, v in row.items()}, '_clean': clean})
    return {'spec': spec, 'rows': checked, 'unknown': unknown, 'name': upload.name,
            'columns': [{'key': k, 'label': label} for k, label, _r, _a in spec['columns'] if k in mapping.values()]}, None


def _summary(rows):
    return {s: sum(1 for r in rows if r['state'] == s) for s in ('ready', 'duplicate', 'error')}


@api_view(['POST'])
@permission_classes([IsSchoolAdmin])
@parser_classes([MultiPartParser, FormParser])
def preview(request, kind):
    result, error = _check(request, kind)
    if error:
        return error
    rows = [{k: v for k, v in r.items() if k != '_clean'} for r in result['rows']]
    return Response({'file': result['name'], 'columns': result['columns'], 'unknown_columns': result['unknown'],
                     'summary': _summary(result['rows']), 'rows': rows})


@api_view(['POST'])
@permission_classes([IsSchoolAdmin])
@parser_classes([MultiPartParser, FormParser])
def commit(request, kind):
    """Adds every ready row in one go (all or nothing); duplicates and rows with errors are skipped and reported."""
    result, error = _check(request, kind)
    if error:
        return error
    spec, rows = result['spec'], result['rows']
    ready = [r for r in rows if r['state'] == 'ready']
    ctx = {'school': request.tenant}
    current = None
    try:
        with transaction.atomic():
            for r in ready:
                current = r['row']
                spec['create'](r['_clean'], ctx)
    except Exception as exc:  # nothing is saved; tell the office which row stopped it
        return Response({'error': f'Nothing was imported: row {current} could not be saved ({str(exc)[:200]}). Fix it and try again.'},
                        status=status.HTTP_400_BAD_REQUEST)
    problems = [{'row': r['row'], 'state': r['state'], 'messages': r['messages'], 'values': r['values']}
                for r in rows if r['state'] != 'ready']
    run = ImportRun.objects.create(tenant=request.tenant, kind=kind, file_name=result['name'][:255], created_by=request.user,
                                   rows=len(rows), added=len(ready), skipped=sum(1 for r in rows if r['state'] == 'duplicate'),
                                   failed=sum(1 for r in rows if r['state'] == 'error'), problems=problems[:1000])
    return Response({'id': str(run.id), 'added': run.added, 'skipped': run.skipped, 'failed': run.failed, 'problems': problems,
                     'message': f'{run.added} added' + (f', {run.skipped} already existed' if run.skipped else '') +
                                (f', {run.failed} had problems' if run.failed else '') + '.'})


@api_view(['GET'])
@permission_classes([IsSchoolAdmin])
def history(request):
    return Response({'results': [{
        'id': str(r.id), 'kind': r.kind, 'label': KINDS.get(r.kind, {}).get('label', r.kind), 'file': r.file_name,
        'when': r.created_at, 'by': (getattr(r.created_by, 'full_name', '') or getattr(r.created_by, 'email', '')) if r.created_by else '',
        'rows': r.rows, 'added': r.added, 'skipped': r.skipped, 'failed': r.failed, 'has_problems': bool(r.problems),
    } for r in ImportRun.objects.select_related('created_by')[:50]]})


@api_view(['GET'])
@permission_classes([IsSchoolAdmin])
def problems_csv(request, run_id):
    run = ImportRun.objects.filter(pk=run_id).first()
    if run is None:
        return Response({'error': 'Not found.'}, status=status.HTTP_404_NOT_FOUND)
    spec = KINDS.get(run.kind, {'columns': []})
    keys = [k for k, *_ in spec['columns']]
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="import-{run.kind}-problems.csv"'
    response.write('﻿')
    writer = csv.writer(response)
    writer.writerow(['Row', 'Problem'] + [label for _, label, *_ in spec['columns']])
    for p in run.problems:
        writer.writerow([p['row'], '; '.join(p['messages'])] + [p['values'].get(k, '') for k in keys])
    return response
