"""Full school export and end-of-contract deletion (P13)."""
import io
import json
import zipfile
from datetime import timedelta
from decimal import Decimal

import pytest
from django.contrib.auth import get_user_model
from django.utils import timezone
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from services.core.audit.models import AuditLog
from services.core.billing import invoicing
from services.core.billing.models import Plan, PlatformInvoice
from services.core.portability.api import run_due
from services.core.portability.models import SchoolDeletion, SchoolExport
from services.core.tenants.context import use_tenant
from services.core.tenants.models import School, TenantMembership
from services.education.academics.models import SchoolClass
from services.education.finance.models import Invoice
from services.education.students.models import Student
from tests.conftest import SchoolFactory, StudentFactory, UserFactory

P = '/api/v1/portability'
User = get_user_model()


def _client(user):
    c = APIClient()
    c.credentials(HTTP_AUTHORIZATION=f'Bearer {RefreshToken.for_user(user).access_token}')
    return c


@pytest.fixture
def port(db, settings, tmp_path):
    settings.MEDIA_ROOT = str(tmp_path)
    s = SchoolFactory(name='Hillside School')
    admin = UserFactory(email='office@hillside.test')
    TenantMembership.objects.create(user=admin, school=s, role='admin', is_primary=True)
    teacher = UserFactory(email='t@hillside.test')
    TenantMembership.objects.create(user=teacher, school=s, role='teacher')
    both = UserFactory(email='shared@both.test')  # works at two schools
    TenantMembership.objects.create(user=both, school=s, role='staff')
    other = SchoolFactory(name='Other School')
    TenantMembership.objects.create(user=both, school=other, role='staff')
    other_admin = UserFactory(email='office@other.test')
    TenantMembership.objects.create(user=other_admin, school=other, role='admin', is_primary=True)
    owner = UserFactory(email='owner@platform.test', is_superuser=True, is_staff=True)
    kw = dict(father_name='', mother_name='', guardian_name='', guardian_phone='')
    with use_tenant(s):
        g1 = SchoolClass.objects.create(tenant=s, name='Grade 1', code='G1')
    amy = StudentFactory(tenant=s, current_class=g1, full_name='Amy Pupil', email='amy@hillside.test', **kw)
    Invoice.objects.create(student=amy, amount=Decimal('500'), due_date=timezone.localdate())
    with use_tenant(other):
        o1 = SchoolClass.objects.create(tenant=other, name='Year 1', code='Y1')
    StudentFactory(tenant=other, current_class=o1, full_name='Elsewhere Kid', **kw)
    return dict(s=s, admin=admin, office=_client(admin), teacher=_client(teacher), both=both, other=other,
                other_office=_client(other_admin), owner=_client(owner), owner_user=owner)


def _read(response):
    body = b''.join(response.streaming_content)
    response.close()  # Windows keeps the file locked until the response is closed
    return body


def _zip(response):
    return zipfile.ZipFile(io.BytesIO(_read(response)))


@pytest.mark.django_db
def test_full_export_in_each_format(port):
    office = port['office']
    r = office.post(f'{P}/exports/', {'format': 'csv'}, format='json')
    assert r.status_code == 201, r.content
    e = r.json()['export']
    assert e['records'] >= 3 and e['people'] >= 3
    z = _zip(office.get(f'{P}/exports/{e["id"]}/download/'))
    names = z.namelist()
    assert {'manifest.json', 'people.csv', 'README.txt', 'education_students.Student.csv', 'education_finance.Invoice.csv'} <= set(names)
    students = z.read('education_students.Student.csv').decode('utf-8-sig')
    assert 'Amy Pupil' in students and 'Elsewhere Kid' not in students  # never another school's records
    assert 'password' not in z.read('people.csv').decode('utf-8-sig').splitlines()[0]
    manifest = json.loads(z.read('manifest.json'))
    assert manifest['school'] == 'Hillside School' and manifest['records']['education_students.Student'] == 1
    # JSON and Excel.
    j = office.post(f'{P}/exports/', {'format': 'json'}, format='json').json()['export']
    data = json.loads(_zip(office.get(f'{P}/exports/{j["id"]}/download/')).read('data.json'))
    assert [s['full_name'] for s in data['education_students.Student']] == ['Amy Pupil'] and data['people']
    x = office.post(f'{P}/exports/', {'format': 'xlsx'}, format='json').json()['export']
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(_read(office.get(f'{P}/exports/{x["id"]}/download/'))))
    assert 'People' in wb.sheetnames and 'Student' in wb.sheetnames
    # Who may: the school's admins (and the platform owner), nobody else.
    assert port['teacher'].post(f'{P}/exports/', {'format': 'csv'}, format='json').status_code == 403
    assert port['other_office'].get(f'{P}/exports/{e["id"]}/download/').status_code == 404
    owner_dl = port['owner'].get(f'{P}/exports/{e["id"]}/download/')
    assert owner_dl.status_code == 200
    _read(owner_dl)
    assert office.post(f'{P}/exports/', {'format': 'pdf'}, format='json').status_code == 400
    # Exports expire after 7 days; the daily job removes the files.
    SchoolExport.objects.filter(pk=e['id']).update(expires_at=timezone.now() - timedelta(minutes=1))
    assert office.get(f'{P}/exports/{e["id"]}/download/').status_code == 410
    assert run_due()['expired_exports'] == 1
    assert SchoolExport.objects.get(pk=e['id']).status == 'expired'
    assert [x['status'] for x in office.get(f'{P}/').json()['exports']].count('expired') == 1


@pytest.mark.django_db
def test_end_of_contract_deletion(port):
    s, office, owner = port['s'], port['office'], port['owner']
    s.refresh_from_db()
    sub_plan = Plan.objects.get(code='standard')
    pinv = invoicing.issue(s, sub_plan, 'monthly', timezone.localdate())
    AuditLog.objects.create(school=s, action='UPDATE', resource_type='v1/x/')
    assert office.post(f'{P}/deletion/', {'confirm': 'Hillside'}, format='json').status_code == 400
    r = office.post(f'{P}/deletion/', {'confirm': 'hillside school', 'reason': 'Contract ended'}, format='json')
    assert r.status_code == 201 and 'You can cancel until then' in r.json()['message']
    d = SchoolDeletion.objects.get(school=s)
    assert (d.scheduled_for - timezone.now()).days >= 29
    assert office.post(f'{P}/deletion/', {'confirm': 'Hillside School'}, format='json').status_code == 400  # already scheduled
    assert office.delete(f'{P}/deletion/').json()['deletion'] is None
    assert SchoolDeletion.objects.get(pk=d.pk).status == 'cancelled'
    office.post(f'{P}/deletion/', {'confirm': 'Hillside School'}, format='json')
    d = SchoolDeletion.objects.get(school=s, status='scheduled')
    assert port['teacher'].post(f'{P}/deletion/', {'confirm': 'Hillside School'}, format='json').status_code == 403
    # Nothing happens before the date; the platform owner can carry it out early with the name typed again.
    assert run_due()['deleted_schools'] == 0
    assert owner.post(f'{P}/platform/{d.pk}/purge/', {'confirm': 'nope'}, format='json').status_code == 400
    r = owner.post(f'{P}/platform/{d.pk}/purge/', {'confirm': 'Hillside School'}, format='json').json()['deletion']
    assert r['status'] == 'done' and r['records_deleted'] >= 3 and r['users_deleted'] >= 3 and r['completed_by'] == 'owner@platform.test'
    assert Student.all_objects.filter(tenant=s).count() == 0 and SchoolClass._base_manager.filter(tenant=s).count() == 0
    assert not Invoice.objects.filter(student__full_name='Amy Pupil').exists() and not AuditLog.objects.filter(school=s).exists()
    assert not User.objects.filter(email__in=['office@hillside.test', 't@hillside.test', 'amy@hillside.test']).exists()
    assert User.objects.filter(email='shared@both.test').exists() and User.objects.filter(email='owner@platform.test').exists()
    school = School.objects.get(pk=s.pk)
    assert not school.is_active and school.name.startswith('Deleted school') and school.settings_json.get('deleted')
    assert PlatformInvoice.objects.filter(pk=pinv.pk).exists()  # financial records stay
    # The other school is untouched.
    assert Student.all_objects.filter(tenant=port['other'], full_name='Elsewhere Kid').exists()
    # Proof, for the platform owner.
    cert = owner.get(f'{P}/deletions/{d.pk}/certificate/').json()['deletion']
    assert cert['counts'].get('education_students.Student') == 1 and cert['school'] == 'Hillside School'
    assert [x['status'] for x in owner.get(f'{P}/platform/').json()['deletions']][:2] == ['done', 'cancelled']


@pytest.mark.django_db
def test_the_daily_job_deletes_on_the_date(port):
    port['office'].post(f'{P}/deletion/', {'confirm': 'Hillside School'}, format='json')
    SchoolDeletion.objects.filter(status='scheduled').update(scheduled_for=timezone.now() - timedelta(minutes=1))
    assert run_due()['deleted_schools'] == 1
    d = SchoolDeletion.objects.get(status='done')
    assert d.completed_by == 'daily job' and Student.all_objects.filter(tenant=port['s']).count() == 0
